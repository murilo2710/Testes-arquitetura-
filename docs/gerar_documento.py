from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
REPORT = DOCS / "arquitetura-e-testes-final"
GITHUB = "https://github.com/murilo2710/Testes-arquitetura-"


def gradle_test_counts() -> tuple[int, int, int]:
    tests = failures = errors = 0
    for report in (ROOT / "build/test-results/test").glob("TEST-*.xml"):
        suite = ElementTree.parse(report).getroot()
        tests += int(suite.get("tests", "0"))
        failures += int(suite.get("failures", "0"))
        errors += int(suite.get("errors", "0"))
    return tests, failures, errors


def build_docx(counts: tuple[int, int, int]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)

    title = document.add_heading("Testes de Arquitetura", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Backend Deviante | Kotlin/Ktor + Python/FastAPI")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Repositório de entrega: " + GITHUB)
    document.add_heading("Integrantes", level=1)
    for name in ("Murilo Regnier", "Bernardo Vieira", "Emanuelle Skolut", "Alander Ávila"):
        document.add_paragraph(name, style="List Bullet")

    document.add_heading("Arquitetura identificada", level=1)
    document.add_paragraph(
        "O BFF é um módulo Gradle Kotlin/Ktor organizado por responsabilidades técnicas: "
        "Routing.kt concentra endpoints e orquestração; repository/ acessa Exposed; db/ declara "
        "tabelas; model/ contém records; dto/ contém contratos HTTP. O microserviço Python separa "
        "a API (main.py), motor (analysis_engine.py), interface/adaptador (drift_detector.py), "
        "algoritmo (ipdd_adwin.py), parser.py e graph_miner.py."
    )
    document.add_paragraph(
        "O estado observado é uma arquitetura em camadas e módulos especializados. Não é uma "
        "migração completa para Clean Architecture/vertical slices: rotas ainda orquestram "
        "diretamente repositórios e parte dos repositórios referencia DTOs. Os testes abaixo "
        "protegem limites reais e não alegam que essa reorganização já foi concluída."
    )

    document.add_heading("Regras automatizadas", level=1)
    document.add_paragraph(
        "ArchUnit (Kotlin/JVM): model não depende de db, DTO, repository, Ktor ou Exposed; "
        "db não depende de DTO, model, repository ou Ktor; DTO não depende da persistência; "
        "repository não depende de Ktor Server."
    )
    document.add_paragraph(
        "PyTestArch (Python): analysis_engine não depende da API, parser, graph_miner ou corpus; "
        "ipdd_adwin não depende do módulo de API. Não há projetos backend Java, C# ou Node neste "
        "checkout; por isso ArchiUnit.Net e ArchiUnitTS não se aplicam."
    )

    document.add_page_break()
    document.add_heading("Comandos e resultados", level=1)
    command_rows = [
        ("Kotlin/Ktor", ".\\gradlew.bat test"),
        ("Python: arquitetura", "cd mining; .\\.venv\\Scripts\\python.exe -m pytest tests/test_architecture.py -q"),
        ("Python: suíte completa", "cd mining; .\\.venv\\Scripts\\python.exe -m pytest tests -q"),
    ]
    commands_table = document.add_table(rows=1, cols=2)
    commands_table.style = "Light Shading Accent 1"
    commands_table.rows[0].cells[0].text = "Execução"
    commands_table.rows[0].cells[1].text = "Comando"
    for label, command in command_rows:
        cells = commands_table.add_row().cells
        cells[0].text = label
        cells[1].text = command
        for run in cells[1].paragraphs[0].runs:
            run.font.name = "Consolas"
            run.font.size = Pt(8)
    document.add_paragraph(
        f"Gradle: BUILD SUCCESSFUL; {counts[0]} testes reportados, "
        f"{counts[1]} falhas e {counts[2]} erros. PyTestArch: 2 passed. "
        "Suíte Python completa: 35 passed, 2 failed."
    )
    document.add_paragraph(
        "As duas falhas existentes dependem dos arquivos de pesquisa "
        "DR_MS_04_matrix.csv e DR_MS_ST_25_matrix.csv, ausentes deste checkout. "
        "A falha reporta ‘Corpus de treinamento ausente: DR_MS_04_matrix.csv.’ "
        "Os dados não foram inventados e os testes não foram contornados."
    )

    document.add_heading("Evidências do código no GitHub", level=1)
    for label, path in (
        ("Teste ArchUnit", "src/test/kotlin/ArchitectureTest.kt"),
        ("Configuração Gradle", "build.gradle.kts"),
        ("Teste PyTestArch", "mining/tests/test_architecture.py"),
        ("Dependências de teste Python", "mining/requirements-dev.txt"),
        ("README de arquitetura", "README.md"),
    ):
        document.add_paragraph(f"{label}: {GITHUB}/blob/main/{path}", style="List Bullet")
    document.add_paragraph(
        "Os links apontam para a branch main. As capturas anexas foram feitas dos relatórios "
        "HTML oficiais gerados pelo Gradle e pelo pytest-html durante a execução."
    )
    document.add_page_break()
    document.add_heading("Captura do relatório Gradle", level=1)
    document.add_paragraph("Relatório oficial: 28 testes, zero falhas, zero ignorados.")
    document.add_picture(str(DOCS / "evidencias/gradle-test-report-recortado.png"), width=Inches(6.2))
    document.add_page_break()
    document.add_heading("Captura do relatório PyTestArch", level=1)
    document.add_paragraph("Relatório oficial do pytest-html: dois testes arquiteturais aprovados.")
    document.add_picture(str(DOCS / "evidencias/python-architecture-test-report-recortado.png"), width=Inches(6.2))
    document.save(REPORT.with_suffix(".docx"))


def build_pdf(counts: tuple[int, int, int]) -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Cover", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#173B57")))
    styles.add(ParagraphStyle(name="BodySmall", parent=styles["BodyText"], fontSize=9, leading=12, spaceAfter=6))
    story = [
        Paragraph("Testes de Arquitetura", styles["Cover"]),
        Paragraph("Backend Deviante | Kotlin/Ktor + Python/FastAPI", styles["Heading2"]),
        Paragraph(f"Repositório: {GITHUB}", styles["BodySmall"]),
        Spacer(1, 6),
        Paragraph("Integrantes", styles["Heading1"]),
        Paragraph("Murilo Regnier<br/>Bernardo Vieira<br/>Emanuelle Skolut<br/>Alander Ávila", styles["BodySmall"]),
        Paragraph("Arquitetura identificada", styles["Heading1"]),
        Paragraph("O BFF é um módulo Kotlin/Ktor organizado por responsabilidades técnicas: Routing.kt concentra endpoints e orquestração; repository/ acessa Exposed; db/ declara tabelas; model/ contém records; dto/ contém contratos HTTP. O microserviço Python separa main.py, analysis_engine.py, drift_detector.py, ipdd_adwin.py, parser.py e graph_miner.py.", styles["BodySmall"]),
        Paragraph("O estado observado é arquitetura em camadas com módulos especializados, não uma migração completa para Clean Architecture/vertical slices. As rotas ainda orquestram repositórios e alguns repositórios referenciam DTOs. Os testes protegem limites reais e não alegam uma reorganização que não ocorreu.", styles["BodySmall"]),
        Paragraph("Regras", styles["Heading1"]),
        Paragraph("ArchUnit: model sem dependência de db/DTO/repository/Ktor/Exposed; db sem dependência de DTO/model/repository/Ktor; DTO sem persistência; repository sem Ktor Server. PyTestArch: analysis_engine sem dependência de API/parser/graph_miner/corpus; ipdd_adwin sem dependência da API. Não há backends Java, C# ou Node neste checkout; ArchiUnit.Net e ArchiUnitTS não se aplicam.", styles["BodySmall"]),
        PageBreak(),
        Paragraph("Execuções verificadas", styles["Heading1"]),
        Paragraph("Kotlin/Ktor: <font name='Courier'>.\\gradlew.bat test</font>", styles["BodySmall"]),
        Paragraph("Python arquitetural: executar de mining com <font name='Courier'>.venv\\Scripts\\python.exe -m pytest tests/test_architecture.py -q</font>.", styles["BodySmall"]),
        Paragraph("Python completo: executar de mining com <font name='Courier'>.venv\\Scripts\\python.exe -m pytest tests -q</font>.", styles["BodySmall"]),
        Spacer(1, 8),
    ]
    data = [
        ["Suíte", "Resultado observado"],
        ["Gradle", f"BUILD SUCCESSFUL; {counts[0]} testes, {counts[1]} falhas, {counts[2]} erros"],
        ["PyTestArch", "2 passed"],
        ["Python completa", "35 passed, 2 failed"],
    ]
    table = Table(data, colWidths=[4 * cm, 12 * cm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173B57")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C7D1")),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F4F7F8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([table, Spacer(1, 7)])
    story.append(Paragraph("As duas falhas Python dependem de DR_MS_04_matrix.csv e DR_MS_ST_25_matrix.csv, ausentes deste checkout. A mensagem observada informa que DR_MS_04_matrix.csv está ausente. Os dados não foram inventados e os testes não foram contornados.", styles["BodySmall"]))
    story.append(Paragraph("Fontes no repositório: src/test/kotlin/ArchitectureTest.kt; build.gradle.kts; mining/tests/test_architecture.py; mining/requirements-dev.txt; README.md. Repositório: " + GITHUB, styles["BodySmall"]))
    story.append(Paragraph("As páginas seguintes contêm capturas dos relatórios HTML oficiais gerados nas execuções.", styles["BodySmall"]))
    story.extend([
        PageBreak(),
        Paragraph("Relatório Gradle", styles["Heading1"]),
        Paragraph("28 testes, zero falhas, zero ignorados; captura do relatório oficial gerado pelo Gradle.", styles["BodySmall"]),
        PdfImage(str(DOCS / "evidencias/gradle-test-report-recortado.png"), width=16 * cm, height=9.95 * cm),
        PageBreak(),
        Paragraph("Relatório PyTestArch", styles["Heading1"]),
        Paragraph("Dois testes de arquitetura aprovados; captura do relatório oficial gerado pelo pytest-html.", styles["BodySmall"]),
        PdfImage(str(DOCS / "evidencias/python-architecture-test-report-recortado.png"), width=16 * cm, height=7.4 * cm),
    ])
    pdf = SimpleDocTemplate(str(REPORT.with_suffix(".pdf")), pagesize=A4, rightMargin=1.6 * cm, leftMargin=1.6 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    pdf.build(story)


if __name__ == "__main__":
    counts = gradle_test_counts()
    if not counts[0] or counts[1] or counts[2]:
        raise SystemExit(f"Relatórios Gradle ausentes ou com falhas: {counts}")
    build_docx(counts)
    build_pdf(counts)
    print(f"Gerados: {REPORT.with_suffix('.docx')} e {REPORT.with_suffix('.pdf')}")
    print(f"Gradle XML: {counts[0]} testes, {counts[1]} falhas, {counts[2]} erros")